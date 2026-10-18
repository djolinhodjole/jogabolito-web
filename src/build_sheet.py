from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
F='Arial'
H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='121A35')
B=Font(name=F); IN=Font(name=F,color='0000FF'); BOLD=Font(name=F,bold=True)
EX=PatternFill('solid',fgColor='EFE7F0')
def head(ws,cols,widths):
    ws.append(cols)
    for i,c in enumerate(ws[1]):
        c.font=H; c.fill=HF; c.alignment=Alignment(horizontal='center')
        ws.column_dimensions[c.column_letter].width=widths[i]
    ws.freeze_panes='A2'
wb=Workbook()
nav=wb.active; nav.title='Navod'
rows=[["Joga Bolíto FC – data pro web"],[""],
["Web čte data z těchto záložek. Názvy záložek a hlaviček NEMĚŇTE, jinak je web nenajde."],[""],
["Po každém zápase:"],
["1. Zapasy: doplňte skóre (Goly JB, Goly souper). Řádek bez skóre = budoucí zápas."],
["2. Statistiky: přidejte řádek pro KAŽDÉHO hráče, který hrál (i s nulami) – z toho se počítají odehrané zápasy."],
["   Sloupce: góly, asistence, nula (1 = vychytaná nula), žluté a červené karty."],
["   Jméno hráče pište přesně jako v záložce Hraci."],
["3. Násobič v Zapasy: 1 = normální odměny, 2 = dvojnásobné (např. zápas proti lídrovi)."],
["4. Odmeny: tady se mění částky za gól, asistenci a vychytanou nulu."],
["5. Tabulka: po každém kole přepište tabulku ligy (pořadí, zápasy, body)."],
["6. Udalosti: každý gól na vlastní řádek – kolo, minuta, tým (JB / Souper), střelec, asistence. Z toho se skládá detail zápasu."],
["7. Prehled: počítá se sám, nic v něm nepřepisujte."],[""],
["Modrý text = ručně zadaná data. Černý text = vzorce."],
["Zdroj dat: zápasové grafiky a výsledky ligy, které dodal Linho (stav po 5. kole, 10.10.2026)."]]
for r in rows: nav.append(r)
nav['A1'].font=Font(name=F,bold=True,size=14)
for r in nav.iter_rows(min_row=2):
    for c in r: c.font=B
nav.column_dimensions['A'].width=100

import json
D=json.load(open('data.json',encoding='utf-8'))
h=wb.create_sheet('Hraci'); head(h,['Jmeno','Cislo','Stav','Foto','Pozice'],[18,8,16,40,14])
roster=[tuple(r) for r in D['roster']]
for r in roster: h.append([r[0],r[1],r[2],None,r[3] or None])
h['E1'].comment=Comment('Brankář / Obránce / Záložník / Útočník. Prázdné = na webu se pozice neukáže.','Claude')
h['D1'].comment=Comment('Odkaz na fotku hráče (na výšku 3:4). Prázdné = na kartě se ukáže erb.','Claude')

z=wb.create_sheet('Zapasy'); head(z,['Kolo','Datum','Cas','Hriste','Doma','Souper','Goly JB','Goly souper','Nasobic'],[7,12,8,10,7,24,9,12,9])
matches=[tuple(None if v=='' else v for v in m) for m in D['matches']]
for m in matches: z.append(list(m))
z['I1'].comment=Comment('Násobič odměn: 1 = normální, 2 = dvojnásobné.','Claude')

s=wb.create_sheet('Statistiky'); head(s,['Kolo','Hrac','Goly','Asistence','Nula','Zlute','Cervene','Nasobic','Odmena','Poznamka'],[7,16,8,11,7,8,9,9,10,16])
st=[tuple(r) for r in D['rows']]
for r in st:
    r=list(r); note=r[7] if len(r)>7 else None; r=r[:7]+[0]*(7-len(r[:7])); s.append(r)
    if note: s.cell(row=s.max_row,column=10,value=note)
s['J1'].comment=Comment('Volitelná poznámka, např. zranění. Ukáže se v detailu hráče.','Claude')
for i in range(2,201):
    s[f'H{i}']=f'=IF(A{i}="","",IFERROR(INDEX(Zapasy!$I$2:$I$200,MATCH(A{i},Zapasy!$A$2:$A$200,0)),1))'
    s[f'I{i}']=f'=IF(B{i}="","",(C{i}*Odmeny!$B$2+D{i}*Odmeny!$B$3+E{i}*Odmeny!$B$4)*IF(H{i}="",1,H{i}))'
    s[f'I{i}'].number_format='#,##0" Kč"'

o=wb.create_sheet('Odmeny'); head(o,['Polozka','Castka'],[28,10])
for r in D['odm']: o.append(list(r))
o['A5'].comment=Comment('Bonusy na konci sezóny – nepočítají se do průběžných rewards.','Claude')
for c in ['B2','B3','B4','B5','B6','B7']: o[c].number_format='#,##0" Kč"'


tb=wb.create_sheet('Tabulka'); head(tb,['Poradi','Tym','Z','V','R','P','Skore','Body'],[8,26,6,6,6,6,9,7])
for r in [(1,'Bílá dáma',5,5,0,0,'36:8',10),(2,'Halandovo kladivo FK',5,4,0,1,'27:11',8),(3,'Kung Funited',5,3,1,1,'18:6',7),(4,'Joga Bolíto',5,3,0,2,'32:11',6),(5,'Milion plus xG',4,2,0,2,'23:9',4),(6,'Příkopy FC',4,2,0,2,'15:10',4),(7,'Kopačky na pivu',4,2,0,2,'12:11',4),(8,'Lobotomici',3,1,1,1,'10:13',3),(9,'Páni pivsonů FC',5,1,1,3,'5:19',3),(10,'Medfúci',4,1,1,2,'5:19',3),(11,'G Yetti Maniacs',4,0,0,4,'2:32',0),(12,'Prague Underground',4,0,0,4,'3:39',0)]:
    tb.append(list(r))
tb['A1'].comment=Comment('Tabulku přepište podle webu ligy po každém kole. Top 3 postupují, poslední 3 sestupují.','Claude')
ud=wb.create_sheet('Udalosti'); head(ud,['Kolo','Minuta','Tym','Strelec','Asistence'],[7,9,10,18,18])
ud['A1'].comment=Comment('Jeden řádek = jeden gól. Tym: JB = náš gól, Souper = gól soupeře. Např.: 6 | 12 | JB | Sultan | Djolinho. U soupeře stačí jméno střelce. Když kolo nemá žádné řádky, web ukáže střelce ze záložky Statistiky.','Claude')
p=wb.create_sheet('Prehled'); head(p,['Hrac','Cislo','Zapasy','Goly','Asistence','Nuly','Zlute','Cervene','Odmena'],[16,8,9,8,11,7,8,9,11])
for i,r in enumerate(roster,start=2):
    p[f'A{i}']=f'=Hraci!A{i}'; p[f'B{i}']=f'=IF(Hraci!B{i}="","",Hraci!B{i})'
    p[f'C{i}']=f'=COUNTIFS(Statistiky!$B$2:$B$200,A{i},Statistiky!$A$2:$A$200,"<>")'
    for col,src in zip('DEFGH','CDEFG'):
        p[f'{col}{i}']=f'=SUMIFS(Statistiky!${src}$2:${src}$200,Statistiky!$B$2:$B$200,A{i})'
    p[f'I{i}']=f'=SUMIFS(Statistiky!$I$2:$I$200,Statistiky!$B$2:$B$200,A{i})'; p[f'I{i}'].number_format='#,##0" Kč"'
t=len(roster)+2
p[f'A{t}']='Celkem'
for col in 'DEFGHI':
    p[f'{col}{t}']=f'=SUM({col}2:{col}{t-1})'
p[f'I{t}'].number_format='#,##0" Kč"'
for ws in wb.worksheets:
    if ws.title=='Navod': continue
    for row in ws.iter_rows(min_row=2):
        for c in row:
            v=c.value
            c.font = B if (isinstance(v,str) and v.startswith('=')) else (IN if v is not None else B)
for c in p[t]: c.font=BOLD
wb.save('JogaBolito_data.xlsx')
